from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\syado\meikoapps\juku_app_report.docx"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_font(run, size=10.5, bold=False, color='000000'):
    run.font.name = 'Meiryo'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Meiryo')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Meiryo')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Meiryo')
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def style_paragraph(p, space_after=5, line=1.12):
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line

def add_para(doc, text='', bold_lead=None, size=10.5, after=5):
    p = doc.add_paragraph()
    style_paragraph(p, after)
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), size, True, '16354D')
        set_font(p.add_run(text[len(bold_lead):]), size)
    else:
        set_font(p.add_run(text), size)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        style_paragraph(p, 2)
        set_font(p.add_run(item), 10)

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_shading(c, '16354D')
        set_cell_border(c)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(h), 9, True, 'FFFFFF')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            if ri % 2 == 1: set_cell_shading(c, 'F3F7F9')
            set_cell_border(c); set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; style_paragraph(p, 0, 1.05)
            set_font(p.add_run(str(val)), 8.8)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62); sec.bottom_margin = Inches(0.58)
sec.left_margin = Inches(0.68); sec.right_margin = Inches(0.68)

styles = doc.styles
styles['Normal'].font.name = 'Meiryo'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Meiryo')
styles['Title'].font.name = 'Meiryo'; styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Meiryo')
styles['Title'].font.size = Pt(22); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
for name, size in [('Heading 1', 15), ('Heading 2', 11.5)]:
    styles[name].font.name = 'Meiryo'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Meiryo')
    styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(title.add_run('生徒管理アプリ 機能 現状 改善点まとめ'), 22, True, '000000')
p = doc.add_paragraph(); style_paragraph(p, 12)
set_font(p.add_run('対象：これまでの実装と要望をもとにした現状整理　作成日：2026年9月20日'), 9.5, False, '5C7281')
add_para(doc, 'この文書は、塾の生徒・授業・講師・シフト・小テストをまとめて管理する現在のアプリについて、できること、実装の到達度、設計上こだわった点、公開運用前に改善すべき点を整理したものです。結論として、画面と業務機能は試作を超えた実用ベータ水準まで広がっていますが、認証・Firestoreの権限・テスト・運用監視は本番公開前に補強が必要です。', size=10.5, after=10)

doc.add_heading('1 アプリの目的と対象業務', level=1)
add_para(doc, '塾の管理者が、日々の授業予定と生徒の所属、講師配置、シフト、長期休みの授業、小テストまでを一つの画面群で扱うためのWebアプリです。PCだけでなくスマートフォンやPWAでも使えること、Firebaseを介して端末間で同じデータを確認できること、必要な帳票をそのまま印刷できることを中心に設計されています。')

doc.add_heading('2 現在できること', level=1)
rows = [
    ('ホーム 本日の授業', '日付移動、コマ別の授業表示、生徒追加、講師割り当て、科目・動画授業表示、当日の小テスト印刷', '実装済み'),
    ('スケジュール', '月間カレンダー、曜日別のB・C・D等の人数、OK・欠表示、休み・長期休み、日付クリックによる授業詳細', '実装済み'),
    ('メンバー', '学年別一覧、検索、選択、個別分析画面、個別分析の印刷、授業内容への導線', '実装済み'),
    ('アカウント', '生徒・非常勤講師の追加・編集・削除、曜日・コマ・科目の固定授業、講師ごとの色', '実装済み'),
    ('長期休み', '期間の作成、対象曜日・コマ設定、利用する生徒の選択、生徒ごとの追加授業、編集・削除', '実装済み'),
    ('科目管理', '学年区分、科目追加・削除、動画科目の設定、授業画面の映表示', '実装済み'),
    ('シフト', '月間シフト編集、自動入力、提出状況、講師別色分け、曜日・祝日の色、月単位削除', '実装済み'),
    ('小テスト・分析', '問題バンク、数学問題、印刷、採点、履歴、正答率・推移・分野別分析', '実装済み'),
    ('Firebase・PWA', 'Firestore保存、端末間同期、localStorageからの移行補助、PWAキャッシュとインストール対応', '実装済み'),
]
add_table(doc, ['領域', '主な機能', '現状'], rows, [1.15, 4.7, 0.8])

doc.add_heading('3 特にこだわってきた点', level=1)
add_bullets(doc, [
    'スマートフォンでの利用を前提に、ヘッダーと下部タブバーを固定し、画面のどこにいても主要操作へ戻れる構成にした。',
    'ホーム、スケジュール、シフトの日付UIを揃え、年・月を個別に選べるようにして、月だけしか選べない状態を改善した。',
    'シフトと小テストの印刷を、画面のプレビュー自身ではなく印刷専用領域から生成する方式にして、PWAやEdgeで印刷ボタンが反応しない問題と同じ内容が二枚出る問題に対応した。',
    'シフト印刷はA4横、小テストはA4縦とし、祝日・土日の色、講師ごとの色、名前の改行防止、日付セル内への収まりを重視した。',
    '長期休みを通常授業と別物にせず、スケジュールやシフト自動入力の対象に含めた。',
    'Firebaseに保存されない、またはスマートフォンだけ表示されない問題を避けるため、Firestoreのデータを優先しつつ、既存localStorageデータを移行する補助も入れた。',
    '日本語の業務用語を見直し、「追加する生徒」から「利用する生徒」へ変更するなど、画面の意味が自然に伝わる表現を採用した。',
])

doc.add_heading('4 現在の完成度の評価', level=1)
add_para(doc, '現時点の位置づけは「機能の広い実用ベータ」です。画面の見た目と主要な業務フローはかなり形になっていますが、認証と権限が仮運用で、複数人が同時に編集する場合の整合性や自動テストが十分ではありません。したがって、実データを扱う内部試用には進めますが、誰でもアクセスできる本番サービスとして公開する前にはセキュリティと運用面の整備が必要です。')
add_table(doc, ['観点', '評価', '理由'], [
    ('画面・操作性', '高め', '主要画面、モバイル、固定ナビ、モーダル、印刷まで一通り実装'),
    ('業務機能の広さ', '高め', '授業・生徒・講師・長期休み・シフト・テストを横断'),
    ('データ保存・同期', '中程度', 'Firestore保存はあるが、端末差・通信失敗・競合の検証が必要'),
    ('認証・権限・安全性', '低め', '仮ログインと全許可ルールが残っており、本番向けではない'),
    ('品質保証・保守性', '中程度', '単一HTML中心で変更は速いが、テストと分割が不足'),
], [1.45, 0.95, 4.25])

doc.add_heading('5 保存とデータの構成', level=1)
add_para(doc, '基本データはFirebase Firestoreに保存する構成です。生徒、講師、利用者、授業割当、曜日別人数、シフト提出、シフト割当、長期休み、科目、小テストといった単位でコレクションまたは設定ドキュメントを分けています。localStorageは表示の一時キャッシュと、旧データをFirestoreへ移すための補助として残っています。')
add_para(doc, 'スマートフォンにシフトが出ない問題では、usersコレクションだけを参照すると非常勤講師が欠落する可能性があったため、accountsとteachersを合わせて表示する修正を入れています。ただし、今後は「どのデータが正本か」を明文化し、同期失敗時の再試行や更新日時の比較まで用意すると、端末差をさらに減らせます。')

doc.add_heading('6 優先して改善すべき点', level=1)
add_table(doc, ['優先度', '改善項目', '理由と対応案'], [
    ('P0', 'Firestoreルールと認証', '現在のfirestore.rulesは全員にread/writeを許可する仮設定。Firebase Authenticationを導入し、管理者・講師・閲覧者などの役割ごとにルールを分ける。'),
    ('P0', '仮ログインの置き換え', 'admin/adminの固定ログインは本番不可。メールリンク、パスワード、またはGoogleログインとパスワード再設定を導入する。'),
    ('P0', 'バックアップと削除保護', '月単位のシフト削除や生徒削除は影響が大きい。確認、復元、変更履歴、定期バックアップを追加する。'),
    ('P1', 'Firestore同期の強化', '通信失敗、同時編集、古いキャッシュを検知し、保存中・保存完了・未保存の状態を明示する。'),
    ('P1', '自動入力のテスト', '通常授業、長期休み、土曜、祝日、月跨ぎ、重複授業を組み合わせたテストデータを自動検証する。'),
    ('P1', '印刷の端末検証', 'Chrome、Edge、PWA、Android、iOSで色、向き、余白、ページ数を確認する。印刷用CSSの回帰テストも用意する。'),
    ('P2', 'コードの分割と型付け', '単一HTMLに機能が集中しているため、画面・Firebase処理・印刷処理を分割し、TypeScriptやテストを導入する。'),
    ('P2', 'アクセシビリティ', 'キーボード操作、フォーカス、ARIA、文字コントラスト、エラー文、タップ領域を点検する。'),
], [0.55, 1.65, 4.45])

doc.add_heading('7 公開前の確認チェックリスト', level=1)
add_bullets(doc, [
    'Firebase Authenticationで実ユーザーがログインでき、ログアウト後に保護画面が開かない。',
    'Firestoreルールを全許可から役割別ルールへ変更し、未ログインの読み書きを拒否できる。',
    'PCとスマートフォンで同じ生徒・授業・シフトが表示され、片方の変更がもう片方へ反映される。',
    '通常授業と長期休みの自動入力で、対象外の日付・曜日・コマが混入しない。',
    'シフト印刷は一回の操作で一部だけが一度印刷され、講師名・色・休日色・年月が正しい。',
    '小テスト印刷は空白ページを作らず、A4縦で問題と記入欄が読みやすい。',
    '削除・月一括削除・生徒削除の前に対象と影響範囲が表示される。',
    'バックアップから復元できることを確認してから実データを投入する。',
])

doc.add_heading('8 まとめ', level=1)
add_para(doc, 'このアプリは、塾の現場で発生する「誰が、いつ、どの授業を担当するか」を中心に、予定・シフト・長期休み・生徒分析・小テストまでを一つにつなげた点が強みです。これまでの改善では、スマートフォンで見やすいこと、印刷物が実務で使えること、Firebaseで端末をまたいで同じ情報を扱えることを一貫して重視してきました。次の段階では、見た目の追加調整よりも、認証・権限・バックアップ・同期・自動テストを優先すると、安心して実運用へ移行できます。')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('生徒管理アプリ 現状整理'), 8, False, '7A8B95')

doc.save(OUT)
print(OUT)
